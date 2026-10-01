"""Prepara un archivo temporal de importación desde el Excel, sin tocar Supabase."""

from collections import defaultdict
from decimal import Decimal
import json
from pathlib import Path
import re

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
PRIVATE = ROOT / ".private_import"
EXCEL = PRIVATE / "excel_inventario.xlsx"
PARSED = PRIVATE / "items_parsed.json"
OUTPUT = PRIVATE / "items_verified.json"
REPORT = ROOT / "reports" / "excel-reconciliation.json"


def clean(value):
    return str(value or "").strip().upper()


def text(value):
    return str(value).strip() if value is not None else ""


def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def source_rows():
    sheet = load_workbook(EXCEL, read_only=True, data_only=True)["INVENTARIO"]
    return [
        {
            "row": number,
            "box": cells[0],
            "name": text(cells[1]),
            "entry": cells[2],
            "unit": text(cells[3]),
            "container": cells[4],
            "outgoing": cells[5],
            "stock": cells[6],
            "condition": text(cells[7]),
        }
        for number, cells in enumerate(sheet.iter_rows(min_col=2, max_col=9, values_only=True), 1)
        if number > 4 and cells[1] is not None
    ]


def score(item, row):
    stock = row["stock"]
    exact = numeric(stock) and Decimal(str(stock)) == Decimal(str(item["cantidad"]))
    truncated = numeric(stock) and int(stock) == item["cantidad"]
    pending = stock == "#VALUE!" and item["cantidad"] == 0
    return (
        3 if exact else 2 if truncated else 1 if pending else 0,
        clean(item["unidad"]) == clean(row["unit"]),
        clean(item.get("especificaciones", {}).get("contenedor")) == clean(row["container"]),
        clean(item.get("especificaciones", {}).get("caja_costal")) == clean(row["box"]),
    )


def use_row(item, row):
    stock = row["stock"]
    pending = not numeric(stock)
    if pending and stock != "#VALUE!":
        raise ValueError(f"Stock no numérico inesperado en fila {row['row']}: {stock!r}")
    if numeric(stock) and (not Decimal(str(stock)).is_finite() or stock < 0):
        raise ValueError(f"Stock inválido en fila {row['row']}: {stock!r}")
    item["cantidad"] = 0 if pending else float(stock) if isinstance(stock, float) else stock
    item["stock_pendiente"] = pending
    item["especificaciones"] = {
        **item.get("especificaciones", {}),
        "fila_excel": row["row"],
        "entrada_excel": text(row["entry"]),
        "salida_excel": text(row["outgoing"]),
        "stock_excel": text(stock),
    }
    return item


def next_code(items, prefix):
    used = [int(m.group(1)) for item in items if (m := re.fullmatch(prefix + r"(\d+)", item["codigo"]))]
    return f"{prefix}{max(used, default=0) + 1:03d}"


def main():
    if not EXCEL.exists() or not PARSED.exists():
        raise SystemExit("Coloque Excel y JSON originales en .private_import antes de preparar la carga.")
    rows = source_rows()
    parsed = json.loads(PARSED.read_text(encoding="utf-8"))
    by_name = defaultdict(list)
    for row in rows:
        by_name[clean(row["name"])].append(row)

    matched = []
    for item in parsed:
        candidates = by_name[clean(item["nombre"])]
        if not candidates:
            raise ValueError(f"No se encontró en Excel el artículo {item['codigo']}: {item['nombre']}")
        best = max(candidates, key=lambda row: score(item, row))
        if score(item, best)[0] == 0:
            raise ValueError(f"Stock incompatible para {item['codigo']} en fila {best['row']}")
        candidates.remove(best)
        matched.append(use_row(item, best))

    missing = [row for group in by_name.values() for row in group]
    added = []
    for row in sorted(missing, key=lambda source: source["row"]):
        category, prefix = ("BATERIAS", "BAT") if "BATERIA" in clean(row["name"]) else ("ESTRUCTURAS", "EST")
        code = next_code(matched + added, prefix)
        item = {
            "id": f"ELM-{len(parsed) + len(added) + 1:04d}",
            "codigo": code,
            "nombre": row["name"],
            "categoria": category,
            "cantidad": 0,
            "stock_minimo": 0,
            "unidad": row["unit"],
            "almacen_id": "ALM-CAL-03",
            "estanteria_id": None,
            "caja_id": None,
            "descripcion": "Ubicación original pendiente de conciliación",
            "especificaciones": {"contenedor": text(row["container"]), "caja_costal": text(row["box"]),
                                  "estado_material": row["condition"]},
        }
        added.append(use_row(item, row))

    verified = matched + added
    if len(verified) != len(rows) or len({item["codigo"] for item in verified}) != len(rows):
        raise ValueError("La conciliación no produjo un artículo único por fila del Excel")
    PRIVATE.mkdir(exist_ok=True)
    OUTPUT.write_text(json.dumps(verified, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    report = {
        "excelRows": len(rows), "matchedParsed": len(matched),
        "addedFromExcelCount": len(added),
        "decimalStockCount": sum(isinstance(item["cantidad"], float) and not item["cantidad"].is_integer()
                                 for item in verified),
        "pendingStockCount": sum(item["stock_pendiente"] for item in verified),
    }
    REPORT.parent.mkdir(exist_ok=True)
    REPORT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Preparados {len(verified)} artículos: {report['decimalStockCount']} stocks decimales, "
          f"{report['pendingStockCount']} pendientes, {len(added)} recuperados del Excel.")


if __name__ == "__main__":
    main()
